"""Lecture des plannings Excel de la serrurerie (VIP Plus).

Deux fichiers comptent pour le pointage :
  · Planning ATE (atelier) — onglet « plan de charge » (le backlog : CH,
    chantier, désignation, heures prévues) et onglet « planning FAB » (qui
    fabrique quoi, jour par jour) ;
  · Planning POSE SER — un onglet par mois, une bande par équipe.

On n'en tire que ce qu'il faut pour pointer : qui est prévu sur quel CH tel
jour (pour proposer « mes chantiers du jour » en tête de liste) et les heures
prévues par CH (pour le point d'affaire). Les couleurs d'Alexis ne sont PAS
lues : elles portent la BU et l'état d'appro, pas le pointage.
"""
import datetime as dt
import re

import openpyxl

from .textes import codes_ch, lignes, normaliser, semaines

# Ce qu'une case de planning contient quand personne ne travaille sur une affaire.
ABSENCES = {"CP", "AM", "FERIE", "FÉRIÉ", "RECUP", "RÉCUP", "MALADIE", "RTT",
            "FORMATION", "ABSENT", "ECOLE", "ÉCOLE"}

MOIS = {"JANVIER": 1, "FEVRIER": 2, "MARS": 3, "AVRIL": 4, "MAI": 5, "JUIN": 6,
        "JUILLET": 7, "AOUT": 8, "SEPTEMBRE": 9, "OCTOBRE": 10, "NOVEMBRE": 11,
        "DECEMBRE": 12}


class FichierInattendu(ValueError):
    """Le fichier déposé n'est pas un planning qu'on sait lire."""


class Grille:
    """Une feuille dont les cellules fusionnées rendent leur valeur partout.

    Dans les plannings, une tâche de trois jours est UNE case fusionnée sur
    trois colonnes : seule la première porte le texte. Lire cellule par
    cellule ferait croire que le gars ne travaille que le lundi."""

    def __init__(self, ws):
        self.ws = ws
        self.origine = {}
        for plage in ws.merged_cells.ranges:
            for r in range(plage.min_row, plage.max_row + 1):
                for c in range(plage.min_col, plage.max_col + 1):
                    self.origine[(r, c)] = (plage.min_row, plage.min_col)

    def valeur(self, r, c):
        r0, c0 = self.origine.get((r, c), (r, c))
        return self.ws.cell(r0, c0).value

    def plage_verticale(self, r, c):
        """(première, dernière) ligne de la fusion qui contient (r, c)."""
        for plage in self.ws.merged_cells.ranges:
            if plage.min_row <= r <= plage.max_row and plage.min_col <= c <= plage.max_col:
                return plage.min_row, plage.max_row
        return r, r


def ouvrir(chemin):
    try:
        return openpyxl.load_workbook(chemin, data_only=True)
    except Exception as e:  # zip corrompu, .xls ancien, autre chose
        raise FichierInattendu(f"Fichier illisible par Excel : {e}") from e


def nature(classeur) -> str:
    """« atelier », « pose » — ou une erreur qui dit ce qu'on a reconnu."""
    noms = [normaliser(n) for n in classeur.sheetnames]
    if "PLANNING FAB" in noms and "PLAN DE CHARGE" in noms:
        entete = normaliser(classeur[classeur.sheetnames[noms.index("PLANNING FAB")]]["A2"].value)
        if "VIP PLUS" in entete or "PLANNING TRAITEMENT" in noms:
            return "atelier"
        raise FichierInattendu("Planning atelier de la menuiserie : prévu pour la V2 "
                               "(la V1 couvre la serrurerie VIP Plus).")
    if any(n.split(" ")[0] in MOIS for n in noms):
        return "pose"
    if any(re.fullmatch(r"20\d\d", n) for n in noms):
        raise FichierInattendu("Planning pose de la menuiserie : prévu pour la V2.")
    if "PLANNING DESSINS" in noms:
        raise FichierInattendu("Planning du bureau d'études : le BET ne pointe pas "
                               "(il est en frais), RHI n'en a pas besoin.")
    raise FichierInattendu("Ni planning atelier, ni planning pose : onglets "
                           + ", ".join(classeur.sheetnames[:6]))


# ═══════════════════════ ATELIER ═══════════════════════

def _dates_colonnes(ws, ligne=2):
    """{colonne: date} de la ligne d'en-tête, erreurs de saisie corrigées.

    Le vrai fichier commence par « 14/12/2024, 15/12/2024, 15/12/2025… » :
    deux années fausses. On lit de droite à gauche et toute date à plus de
    trois jours de « la suivante moins un » est remplacée par celle-ci."""
    brutes = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(ligne, c).value
        if isinstance(v, dt.datetime):
            brutes[c] = v.date()
        elif isinstance(v, dt.date):
            brutes[c] = v
    colonnes = sorted(brutes)
    propres = {}
    suivante = None
    for c in reversed(colonnes):
        d = brutes[c]
        if suivante is not None:
            attendue = suivante - dt.timedelta(days=1)
            if abs((d - attendue).days) > 3:
                d = attendue
        propres[c] = d
        suivante = d
    return propres


def _nom_personne(v) -> str:
    n = " ".join(str(v).split()).strip()
    return n.upper()


def lire_atelier(classeur) -> dict:
    """Personnes de l'atelier, leurs affectations par jour, et le backlog."""
    noms = {normaliser(n): n for n in classeur.sheetnames}
    fab = classeur[noms["PLANNING FAB"]]
    grille = Grille(fab)
    dates = _dates_colonnes(fab)

    # Les noms des opérateurs sont écrits dans la colonne du samedi, sur une
    # case fusionnée qui couvre toute leur bande (libellé, commentaire, CH,
    # responsable). On relit chaque samedi : l'équipe bouge dans l'année.
    bandes = {}   # (premiere_ligne, derniere_ligne) -> nom, au samedi le plus récent
    for c, d in sorted(dates.items()):
        if d.weekday() != 5:
            continue
        for r in range(3, min(fab.max_row, 80) + 1):
            if (r, c) in grille.origine and grille.origine[(r, c)] != (r, c):
                continue
            v = fab.cell(r, c).value
            if not v or not isinstance(v, str):
                continue
            nom = _nom_personne(v)
            if nom in ABSENCES or codes_ch(nom) or "SOUS" in nom or len(nom) > 30:
                continue
            bandes[grille.plage_verticale(r, c)] = nom

    affectations = []
    for (r0, r1), nom in sorted(bandes.items()):
        for c, jour in dates.items():
            if jour.weekday() >= 5:
                continue
            codes, libelle = [], None
            for r in range(r0, r1 + 1):
                v = grille.valeur(r, c)
                trouves = codes_ch(v)
                if trouves:
                    codes += [x for x in trouves if x not in codes]
                elif libelle is None and v and r == r0:
                    libelle = " / ".join(lignes(v))
            if libelle and normaliser(libelle) in ABSENCES:
                continue
            if codes or libelle:
                affectations.append({"personne": nom, "jour": jour,
                                     "codes": codes, "libelle": libelle or ""})

    return {"personnes": sorted(set(bandes.values())),
            "affectations": affectations,
            "affaires": lire_backlog(classeur[noms["PLAN DE CHARGE"]])}


def lire_backlog(ws) -> list:
    """Le plan de charge : une ligne par pièce à fabriquer.

    Colonnes repérées par leur titre (ligne « N° Affaire ») et non par leur
    lettre : Alexis ajoute des colonnes, un index fixe se décalerait en
    silence."""
    entete, cols = None, {}
    for r in range(1, 15):
        titres = {normaliser(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)}
        if "N AFFAIRE" in titres:
            entete, cols = r, titres
            break
    if entete is None:
        raise FichierInattendu("Plan de charge : ligne « N° Affaire » introuvable.")

    def col(*noms):
        for n in noms:
            if n in cols:
                return cols[n]
        return None

    c_ch, c_aff = col("N AFFAIRE"), col("AFFAIRE")
    c_des, c_cond = col("DESIGNATION"), col("CONDUC")
    c_h, c_sem, c_com = col("PREVISIONNEL H H"), col("S FAB"), col("COMMENTAIRE")

    lignes_backlog = []
    for r in range(entete + 1, ws.max_row + 1):
        chantier = ws.cell(r, c_aff).value if c_aff else None
        designation = ws.cell(r, c_des).value if c_des else None
        if not chantier and not designation:
            continue
        if designation and "TOTAL" in normaliser(designation):
            continue
        heures = ws.cell(r, c_h).value if c_h else None
        try:
            heures = float(heures) if heures not in (None, "") else None
        except (TypeError, ValueError):
            heures = None
        codes = codes_ch(ws.cell(r, c_ch).value) if c_ch else []
        lignes_backlog.append({
            "ch": codes[0] if codes else None,
            "chantier": " ".join(str(chantier or "").split()),
            "designation": " ".join(str(designation or "").split()),
            "conduc": str(ws.cell(r, c_cond).value or "").strip() if c_cond else "",
            "heures": heures,
            "semaines": semaines(ws.cell(r, c_sem).value) if c_sem else [],
            "commentaire": str(ws.cell(r, c_com).value or "").strip() if c_com else "",
        })
    return lignes_backlog


# ═══════════════════════ POSE ═══════════════════════

SOUS_LIGNES = {"N AFFAIRE", "POSE IMPERATIVE", "CONDUC", "ABSCENCE TECH",
               "ABSENCE TECH", "INTERIM REMPLACEMENT"}


def personnes_equipe(libelle) -> list:
    """« EQUIPE LUC & MARC\\n(Resp. MM) » → ["LUC", "MARC"]."""
    premiere = lignes(libelle)[0] if lignes(libelle) else ""
    t = normaliser(premiere)
    t = re.sub(r"^EQUIPE\s+", "", t)
    morceaux = re.split(r"\s+(?:ET)\s+|\s*&\s*", premiere.upper().replace("EQUIPE", ""))
    noms = []
    for m in morceaux:
        m = normaliser(m)
        if m and "SOUS TRAITANT" not in m and m not in noms:
            noms.append(m)
    return noms if t else []


def lire_pose(classeur) -> dict:
    """Équipes de pose, leurs membres, et l'affectation jour par jour."""
    personnes, affectations = set(), []
    for nom_onglet in classeur.sheetnames:
        mots = normaliser(nom_onglet).split()
        if len(mots) != 2 or mots[0] not in MOIS or not mots[1].isdigit():
            continue
        mois, annee = MOIS[mots[0]], int(mots[1])
        ws = classeur[nom_onglet]
        grille = Grille(ws)

        jours = {}
        for c in range(2, ws.max_column + 1):
            v = ws.cell(2, c).value
            try:
                jours[c] = dt.date(annee, mois, int(v))
            except (TypeError, ValueError):
                continue

        for r in range(3, ws.max_row + 1):
            etiquette = ws.cell(r, 1).value
            if not etiquette or normaliser(etiquette) in SOUS_LIGNES:
                continue
            suite = [normaliser(ws.cell(r + k, 1).value) for k in (1, 2)]
            if "N AFFAIRE" not in suite and "POSE IMPERATIVE" not in suite:
                continue   # pas une bande d'équipe (encadrants, légende…)
            membres = personnes_equipe(etiquette)
            if not membres:
                continue   # sous-traitant poseur : il ne pointe pas chez nous
            personnes.update(membres)
            r_ch = r + 1 if suite[0] == "N AFFAIRE" else None
            for c, jour in jours.items():
                v = grille.valeur(r, c)
                if not v or normaliser(v) in ABSENCES:
                    continue
                codes = codes_ch(grille.valeur(r_ch, c)) if r_ch else []
                codes += [x for x in codes_ch(v) if x not in codes]
                for m in membres:
                    affectations.append({"personne": m, "jour": jour, "codes": codes,
                                         "libelle": " / ".join(lignes(v))})
    return {"personnes": sorted(personnes), "affectations": affectations}


def lire(chemin) -> dict:
    """Point d'entrée : reconnaît le fichier et le lit."""
    classeur = ouvrir(chemin)
    genre = nature(classeur)
    donnees = lire_atelier(classeur) if genre == "atelier" else lire_pose(classeur)
    donnees["nature"] = genre
    return donnees
