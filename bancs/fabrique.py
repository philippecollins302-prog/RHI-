"""Fabrique de plannings FICTIFS, construits comme ceux d'Alexis.

Le dépôt est public : les vrais plannings (noms de clients, d'ouvriers,
numéros d'affaire) n'y entrent pas. Ces classeurs reproduisent la forme —
cases fusionnées, noms au samedi, erreurs de date dans l'en-tête, bande
« N° AFFAIRE » sous chaque équipe — avec des noms inventés.
"""
import datetime as dt

import openpyxl


def atelier(chemin):
    wb = openpyxl.Workbook()
    pc = wb.active
    pc.title = "plan de charge "
    pc["B1"] = "CHARGE ATELIER"
    pc.append([])
    pc.append([])
    pc["A4"] = "EN PROVISION"
    for i, t in enumerate(["N° Affaire", "Affaire", "Désignation", "Conduc.",
                           "Prévisionnel h/h", "S - FAB", "COMMENTAIRE"], start=1):
        pc.cell(5, i, t)
    lignes = [
        ("CH00901", "LES PINS", "GARDE-CORPS (X4)", "AA", 40, "S40", ""),
        ("CH00901", "LES PINS", "MAIN COURANTE 12ML", "AA", 16, "S41", ""),
        ("CH00902", "LE PORT", "PORTAIL COULISSANT", "BB", 54, "S40-41", "ATTENTE VALIDATION"),
        (None, "SANS NUMERO", "POTELET", "BB", 4, "S39", ""),
    ]
    for r, l in enumerate(lignes, start=6):
        for c, v in enumerate(l, start=1):
            pc.cell(r, c, v)
    pc.cell(6 + len(lignes), 3, "TOTAL HEURES - CHARGE ATELIER")
    pc.cell(6 + len(lignes), 5, 114)

    fab = wb.create_sheet("planning FAB")
    fab["A2"] = "VIP PLUS\n\n-\nPLANNING PREVISIONNEL"
    # Semaine 40 de 2026 : lundi 28/09 → dimanche 04/10, colonnes B..H.
    lundi = dt.date(2026, 9, 28)
    for i in range(7):
        fab.cell(2, 2 + i, dt.datetime.combine(lundi + dt.timedelta(days=i), dt.time()))
    # L'erreur du vrai fichier : une année fausse en tête de ligne.
    fab.cell(2, 2, dt.datetime(2024, 9, 28))
    wb.create_sheet("planning TRAITEMENT")

    def bande(r0, nom, taches):
        fab.cell(r0, 7, nom)                       # samedi = colonne G
        fab.merge_cells(start_row=r0, start_column=7, end_row=r0 + 3, end_column=8)
        for col, (libelle, ch, largeur) in taches.items():
            fab.cell(r0, col, libelle)
            fab.cell(r0 + 2, col, ch)
            fab.cell(r0 + 3, col, "AA")
            if largeur > 1:
                for rr in (r0, r0 + 2, r0 + 3):
                    fab.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + largeur - 1)

    bande(3, "PAUL", {2: ("LES PINS\nGARDE-CORPS (X4)", "CH00901", 3),
                      5: ("LE PORT\nPORTAIL COULISSANT", "CH00902", 2)})
    bande(7, "JEAN", {2: ("CP", None, 1), 3: ("LE PORT\nPORTAIL", "CH00902/CH00901", 1)})
    fab.cell(11, 7, "SOUS TRAITANT - TOTO")
    wb.save(chemin)


def pose(chemin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SEPTEMBRE 2026"
    ws["A1"] = "SEPTEMBRE 2026"
    for j in range(1, 31):
        ws.cell(2, 1 + j, j)
    ws["A4"] = "EQUIPE LUC & MARC\n(Resp. AA)"
    ws["A5"] = "N° AFFAIRE"
    ws["A6"] = "POSE IMPERATIVE"
    ws["A7"] = "Conduc."
    ws.cell(4, 29, "LES PINS\nGARDE-CORPS")      # 28/09, sur deux jours
    ws.merge_cells(start_row=4, start_column=29, end_row=4, end_column=30)
    ws.cell(5, 29, "CH00901")
    ws.merge_cells(start_row=5, start_column=29, end_row=5, end_column=30)
    ws.cell(4, 31, "CP")                         # 30/09
    ws["A10"] = "SOUS-TRAITANT POSEUR"
    ws["A11"] = "POSE IMPERATIVE"
    ws.cell(10, 29, "LE PORT\nPEINTURE")
    ws["A14"] = "ALEXIS"                        # encadrant : pas une équipe
    ws.cell(14, 29, "AM")
    wb.save(chemin)


def bet(chemin):
    wb = openpyxl.Workbook()
    wb.active.title = "Plan de charge"
    wb.create_sheet("Planning Dessins")
    wb.save(chemin)


def atelier_men(chemin):
    """Planning atelier menuiserie : noms en colonne B, lignes « N° AFFAIRES »."""
    wb = openpyxl.Workbook()
    pc = wb.active
    pc.title = "Plan de charge"
    pc["A3"], pc["B3"] = "Resp.", "Désignation"
    pc["A5"] = "Resp1"
    pc["B6"] = "Ecole des Oliviers CH00911"
    fab = wb.create_sheet("Planning FAB")
    lundi = dt.date(2026, 9, 28)
    for i in range(7):
        fab.cell(3, 3 + i, dt.datetime.combine(lundi + dt.timedelta(days=i), dt.time()))
    fab["A4"], fab["B4"] = "AA", "/APPRO VERRE"       # chargés d'affaires : pas des opérateurs
    fab["A10"], fab["B10"] = "DEBIT", "HUGO"
    fab.cell(10, 3, "ECOLE DES OLIVIERS")
    fab.merge_cells(start_row=10, start_column=3, end_row=10, end_column=4)
    fab.cell(11, 3, "CH00911")
    fab.merge_cells(start_row=11, start_column=3, end_row=11, end_column=4)
    fab["A13"], fab["B13"] = "FABRICATION", "NINA"
    fab["B14"], fab["B15"] = "N° AFFAIRES", "RA"
    fab.cell(13, 5, "SANS NUMERO ENCORE")          # pas de CH : affectation sans affaire
    fab.cell(13, 6, "cp")
    fab["B17"] = "DATE /POSE/DEMANDE PAR RA"       # titre du bloc suivant
    wb.save(chemin)


def pose_men(chemin):
    """Planning pose menuiserie : onglet par année, « L28 » en ligne 3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.cell(1, 2, "SEPTEMBRE")
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
    ws.cell(1, 5, "OCTOBRE")
    for c, j in ((2, "L28"), (3, "M29"), (4, "M30"), (5, "J1")):
        ws.cell(3, c, j)
    ws["A3"] = "EQUIPES"
    ws["A4"] = "OSCAR \nPAUL"                      # membres séparés par un retour à la ligne
    ws["A6"], ws["A7"] = "N° AFFAIRE", "CA"
    ws.cell(4, 2, "ECOLE DES OLIVIERS\nMENUISERIES")
    ws.cell(6, 2, "CH00911")
    ws.cell(4, 3, "CP")
    ws.cell(5, 4, "SAV DIVERS")                    # deuxième ligne de la bande
    ws["A9"] = "ST SOUS-TRAITANT"                  # sous-traitant : ne pointe pas
    ws["A10"], ws["A11"] = "N° AFFAIRE", "CA"
    ws.cell(9, 2, "CHANTIER X")
    wb.save(chemin)
