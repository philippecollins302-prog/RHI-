"""Les plannings relus tout seuls (revue avec Alexis, 01/10/2026).

« À chaque fois qu'on change le planning, il faut aller le redéposer ;
l'idéal, une boucle toutes les cinq minutes. » RHI relit des liens de
téléchargement : il importe un fichier changé, ne refait rien d'un fichier
inchangé, explique un lien qui donne une page web, survit à un lien mort —
et n'écrit jamais le lien (il porte un jeton)."""
import os
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient

from bancs import fabrique
from bancs.outils import fin, verif

tmp = Path(tempfile.mkdtemp())
os.environ["RHI_DONNEES"] = str(tmp)
for k in ("RHI_CODE_TERRAIN", "RHI_CODE_BUREAU", "INTERFAST_VIP"):
    os.environ.pop(k, None)
JETON = "jeton-secret-0123456789"
ATE = f"https://partage.exemple.fr/ate.xlsx?download=1&token={JETON}"
POSE = f"https://partage.exemple.fr/pose.xlsx?download=1&token={JETON}x"
MORT = f"https://mort.exemple.fr/x?token={JETON}y"
os.environ["RHI_PLANNINGS_URL"] = f"{ATE}\n{POSE}  {MORT} http://pas-chiffre.fr/z"

import app as appli  # noqa: E402
from rhi import base, depot_auto  # noqa: E402

appli.app.state.chemin_base = tmp / "rhi.db"
c = TestClient(appli.app)
db = base.connexion(tmp / "rhi.db")
fabrique.atelier(tmp / "ate.xlsx")
fabrique.pose(tmp / "pose.xlsx")
fichiers = {ATE: (tmp / "ate.xlsx").read_bytes(), POSE: (tmp / "pose.xlsx").read_bytes()}
lus = []


def lien(url):
    lus.append(url)
    if url not in fichiers:
        raise ConnectionError(f"impossible de joindre {url}")
    return fichiers[url]


verif(depot_auto.sources() == [ATE, POSE, MORT], "trois liens https ; le lien en clair (http) est écarté")
r = {x["source"]: x for x in depot_auto.relever(db, "VIP", lien)}
ate, pose, mort = (depot_auto.etiquette(u) for u in (ATE, POSE, MORT))
verif(r[ate]["statut"].startswith("importé") and r[ate]["nature"] == "atelier", f"l'atelier importé : {r[ate]}")
verif(r[pose]["statut"].startswith("importé") and r[pose]["nature"] == "pose", f"la pose importée : {r[pose]}")
verif(r[mort]["statut"].startswith("injoignable"), f"le lien mort ne bloque pas les autres : {r[mort]}")
verif(db.execute("SELECT COUNT(*) FROM planning").fetchone()[0] > 0, "le planning est sur les tablettes")

# Rien n'a changé : rien n'est réimporté.
avant = db.execute("SELECT COUNT(*) FROM planning").fetchone()[0]
r = {x["source"]: x for x in depot_auto.relever(db, "VIP", lien)}
verif(r[ate]["statut"] == "inchangé" and r[pose]["statut"] == "inchangé", "fichiers inchangés : rien à refaire")

# Alexis enregistre le planning : JEAN passe sur le CH00902 lundi.
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(tmp / "ate.xlsx")
wb["planning FAB"].cell(7, 2, "LE PORT\nPORTAIL")
wb["planning FAB"].cell(9, 2, "CH00902")
wb.save(tmp / "ate2.xlsx")
fichiers[ATE] = (tmp / "ate2.xlsx").read_bytes()
r = {x["source"]: x for x in depot_auto.relever(db, "VIP", lien)}
verif(r[ate]["statut"].startswith("importé"), f"le planning modifié est réimporté : {r[ate]}")
verif(db.execute("SELECT 1 FROM planning WHERE personne='JEAN' AND jour='2026-09-28' AND ch='CH00902'").fetchone(),
      "la modification d'Alexis est arrivée sans redépôt")

# Un lien de partage qui donne la page web au lieu du fichier : dit comment faire.
fichiers[POSE] = b"<!DOCTYPE html><html>Connexion</html>"
r = {x["source"]: x for x in depot_auto.relever(db, "VIP", lien)}
verif("download=1" in r[pose]["statut"], f"page web au lieu du fichier : le remède est dit : {r[pose]}")
verif(db.execute("SELECT COUNT(*) FROM planning WHERE origine='pose'").fetchone()[0] > 0,
      "un mauvais passage ne vide pas le planning en place")

# Le bureau voit l'état ; le jeton n'est écrit nulle part.
os.environ.pop("RHI_CODE_BUREAU", None)
e = c.get("/api/plannings/auto").json()
verif(e["actif"] and e["minutes"] == 5 and len(e["fichiers"]) == 3, f"le bureau voit les trois fichiers : {e}")
tout = str(e) + str([tuple(x) for x in db.execute("SELECT * FROM depots_auto")])
verif(JETON not in tout and "token" not in tout, "le lien (et son jeton) n'est ni en base ni à l'écran")
verif(all(x["sha"] is None for x in e["fichiers"]), "l'empreinte interne n'est pas montrée")

# Sans lien réglé : rien ne tourne, le bureau dit quoi faire.
os.environ["RHI_PLANNINGS_URL"] = ""
verif(depot_auto.relever(db, "VIP", lien) == [] and not c.get("/api/plannings/auto").json()["actif"],
      "sans RHI_PLANNINGS_URL : aucun passage")
os.environ["RHI_PLANNINGS_MINUTES"] = "zéro"
verif(depot_auto.minutes() == 5, "un réglage de période illisible : cinq minutes")
os.environ.pop("RHI_PLANNINGS_MINUTES")
os.environ.pop("RHI_PLANNINGS_URL")

fin("banc-depot-auto")
